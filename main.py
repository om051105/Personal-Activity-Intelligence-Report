import os
import glob
import openpyxl
from statistics import mean

# Find the first .xlsx file in the current directory
excel_files = glob.glob("*.xlsx")
if not excel_files:
    raise FileNotFoundError("No .xlsx file found in the current directory.")

file_name = excel_files[0]
print("Using file:", file_name)

print("Libraries imported successfully")
# ============================================================
# CAP776 - MY DATA, MY STORY
# ============================================================


# ------------------------------------------------------------
# 1. OPEN EXCEL FILE
# ------------------------------------------------------------

try:

    workbook = openpyxl.load_workbook(
        file_name,
        data_only=True
    )

    # Use Daily Log if available
    if "Daily Log" in workbook.sheetnames:
        sheet = workbook["Daily Log"]
    else:
        sheet = workbook[workbook.sheetnames[0]]

    print("Excel file opened successfully")
    print("Sheet:", sheet.title)

except Exception as error:

    print("Error opening Excel file:")
    print(error)

    raise


# ------------------------------------------------------------
# 2. FIND HEADER ROW
# ------------------------------------------------------------

def find_header_row(sheet):

    for row in range(1, sheet.max_row + 1):

        for column in range(1, sheet.max_column + 1):

            if sheet.cell(row, column).value == "Date":

                return row

    return None


header_row = find_header_row(sheet)

if header_row is None:

    raise ValueError(
        "Date column was not found."
    )

print("Header row:", header_row)


# ------------------------------------------------------------
# 3. READ HEADERS
# ------------------------------------------------------------

headers = []

for column in range(
    1,
    sheet.max_column + 1
):

    headers.append(
        sheet.cell(
            header_row,
            column
        ).value
    )


# ------------------------------------------------------------
# 4. CHECK REQUIRED COLUMNS
# ------------------------------------------------------------

required_columns = [

    "Date",
    "Sleep (min)",
    "Fitness (min)",
    "Study (min)",
    "Coding (min)",
    "Class (min)",
    "Other Activities (min)",
    "Total Tracked (min)",
    "Free/Unaccounted (min)",
    "Day's Feeling",
    "Satisfaction Level",
    "Energy Level"

]


for column in required_columns:

    if column not in headers:

        raise ValueError(
            "Missing column: " + column
        )


# ------------------------------------------------------------
# 5. READ DATA
# ------------------------------------------------------------

records = []

for row in range(
    header_row + 1,
    sheet.max_row + 1
):

    values = []

    for column in range(
        1,
        len(headers) + 1
    ):

        values.append(
            sheet.cell(
                row,
                column
            ).value
        )


    if all(
        value is None
        for value in values
    ):

        continue


    record = dict(
        zip(
            headers,
            values
        )
    )


    records.append(record)


print("Total rows:", len(records))


# ------------------------------------------------------------
# 6. SCORING FUNCTIONS
# ------------------------------------------------------------

def feeling_score(value):

    scores = {

        "Excellent": 5,
        "Good": 4,
        "Neutral": 3,
        "Low": 2,
        "Stressed": 1

    }

    return scores.get(
        str(value).strip(),
        0
    )


def satisfaction_score(value):

    scores = {

        "Very Satisfied": 5,
        "Satisfied": 4,
        "Neutral": 3,
        "Unsatisfied": 2,
        "Very Unsatisfied": 1

    }

    return scores.get(
        str(value).strip(),
        0
    )


def energy_score(value):

    scores = {

        "High": 3,
        "Medium": 2,
        "Low": 1

    }

    return scores.get(
        str(value).strip(),
        0
    )


# ------------------------------------------------------------
# 7. VALIDATE DATA
# ------------------------------------------------------------

def validate_record(record):

    try:

        if record["Date"] is None:

            return False


        columns = [

            "Sleep (min)",
            "Fitness (min)",
            "Study (min)",
            "Coding (min)",
            "Class (min)",
            "Other Activities (min)",
            "Total Tracked (min)",
            "Free/Unaccounted (min)"

        ]


        for column in columns:

            value = record[column]

            if value is None:
                return False

            if not isinstance(
                value,
                (int, float)
            ):
                return False

            if value < 0:
                return False


        if feeling_score(
            record["Day's Feeling"]
        ) == 0:

            return False


        if satisfaction_score(
            record["Satisfaction Level"]
        ) == 0:

            return False


        if energy_score(
            record["Energy Level"]
        ) == 0:

            return False


        # Check total tracked time

        total = (

            record["Sleep (min)"]
            +
            record["Fitness (min)"]
            +
            record["Study (min)"]
            +
            record["Coding (min)"]
            +
            record["Class (min)"]
            +
            record["Other Activities (min)"]

        )


        if abs(
            total -
            record["Total Tracked (min)"]
        ) > 0.01:

            return False


        # Check free time

        free_time = (

            1440 -
            record["Total Tracked (min)"]

        )


        if abs(
            free_time -
            record["Free/Unaccounted (min)"]
        ) > 0.01:

            return False


        return True


    except:

        return False


# ------------------------------------------------------------
# 8. VALID AND INVALID DATA
# ------------------------------------------------------------

valid_records = []
invalid_records = []


for record in records:

    if validate_record(record):

        valid_records.append(record)

    else:

        invalid_records.append(record)


print("Valid records:", len(valid_records))
print("Invalid records:", len(invalid_records))


if len(valid_records) == 0:

    raise ValueError(
        "No valid records found."
    )


# ------------------------------------------------------------
# 9. AVERAGE FUNCTION
# ------------------------------------------------------------

def average(records, column):

    values = []

    for record in records:

        values.append(
            record[column]
        )

    return mean(values)


# ------------------------------------------------------------
# 10. DAILY AVERAGES
# ------------------------------------------------------------

average_sleep = average(
    valid_records,
    "Sleep (min)"
)

average_fitness = average(
    valid_records,
    "Fitness (min)"
)

average_study = average(
    valid_records,
    "Study (min)"
)

average_coding = average(
    valid_records,
    "Coding (min)"
)

average_class = average(
    valid_records,
    "Class (min)"
)

average_other = average(
    valid_records,
    "Other Activities (min)"
)

average_total = average(
    valid_records,
    "Total Tracked (min)"
)

average_free = average(
    valid_records,
    "Free/Unaccounted (min)"
)


# ------------------------------------------------------------
# 11. EXPERIENCE INDEX
# ------------------------------------------------------------

experience_scores = []


for record in valid_records:

    feeling = feeling_score(
        record["Day's Feeling"]
    )

    satisfaction = satisfaction_score(
        record["Satisfaction Level"]
    )

    energy = energy_score(
        record["Energy Level"]
    )


    score = (

        feeling
        +
        satisfaction
        +
        energy

    ) / 3


    experience_scores.append(score)


EI = mean(
    experience_scores
)


# ------------------------------------------------------------
# 12. REQUIRED INDICES
# ------------------------------------------------------------

# Tech Productivity Index
TPI = average_coding


# Academic Activity Index
AAI = (
    average_study
    +
    average_class
)


# Physical Activity Index
PhAI = average_fitness


# Sleep and Recovery Index
SRI = average_sleep


# Activity Balance Index
ABI = average_free


# Time Utilization Index
TUI = average_total


# ------------------------------------------------------------
# 13. DATA CONTINUITY INDEX
# ------------------------------------------------------------

EXPECTED_DAYS = 40

DCI = (

    len(valid_records)
    /
    EXPECTED_DAYS

) * 100


# ------------------------------------------------------------
# 14. PERSONAL ACTIVITY INDEX
# ------------------------------------------------------------

PAI = (

    0.15 * TPI
    +
    0.20 * AAI
    +
    0.15 * PhAI
    +
    0.20 * SRI
    +
    0.15 * TUI
    +
    0.10 * EI
    +
    0.05 * DCI

)


# ------------------------------------------------------------
# 15. CORRELATION FUNCTION
# ------------------------------------------------------------

def correlation(x, y):

    x_mean = mean(x)
    y_mean = mean(y)

    numerator = 0
    x_total = 0
    y_total = 0


    for a, b in zip(x, y):

        numerator += (
            (a - x_mean)
            *
            (b - y_mean)
        )

        x_total += (
            (a - x_mean) ** 2
        )

        y_total += (
            (b - y_mean) ** 2
        )


    denominator = (
        x_total * y_total
    ) ** 0.5


    if denominator == 0:

        return 0


    return numerator / denominator


# ------------------------------------------------------------
# 16. SLEEP - ENERGY
# ------------------------------------------------------------

sleep_values = []
energy_values = []


for record in valid_records:

    sleep_values.append(
        record["Sleep (min)"]
    )

    energy_values.append(
        energy_score(
            record["Energy Level"]
        )
    )


sleep_energy = correlation(
    sleep_values,
    energy_values
)


# ------------------------------------------------------------
# 17. STUDY - SATISFACTION
# ------------------------------------------------------------

study_values = []
satisfaction_values = []


for record in valid_records:

    study_values.append(
        record["Study (min)"]
    )

    satisfaction_values.append(
        satisfaction_score(
            record["Satisfaction Level"]
        )
    )


study_satisfaction = correlation(
    study_values,
    satisfaction_values
)


# ------------------------------------------------------------
# 18. CODING - ENERGY
# ------------------------------------------------------------

coding_values = []


for record in valid_records:

    coding_values.append(
        record["Coding (min)"]
    )


coding_energy = correlation(
    coding_values,
    energy_values
)


# ------------------------------------------------------------
# 19. INTERPRET CORRELATION
# ------------------------------------------------------------

def interpretation(value):

    if value >= 0.70:

        return "Strong positive relationship"

    elif value >= 0.40:

        return "Moderate positive relationship"

    elif value >= 0.20:

        return "Weak positive relationship"

    elif value <= -0.70:

        return "Strong negative relationship"

    elif value <= -0.40:

        return "Moderate negative relationship"

    elif value <= -0.20:

        return "Weak negative relationship"

    else:

        return "Very weak or no clear relationship"


# ============================================================
# 20. FINAL OUTPUT
# ============================================================

print()
print("=" * 60)

print("CAP776 - MY DATA, MY STORY")

print("PERSONAL ACTIVITY INTELLIGENCE REPORT")

print("=" * 60)


# ------------------------------------------------------------
# SUMMARY
# ------------------------------------------------------------

print()
print("DATA SUMMARY")
print("-" * 60)

print(
    "Expected days:",
    EXPECTED_DAYS
)

print(
    "Valid days:",
    len(valid_records)
)

print(
    "Missing days:",
    EXPECTED_DAYS - len(valid_records)
)

print(
    "Invalid records:",
    len(invalid_records)
)


# ------------------------------------------------------------
# AVERAGES
# ------------------------------------------------------------

print()
print("AVERAGE DAILY ACTIVITY")
print("-" * 60)

print(
    "Sleep:",
    round(average_sleep, 2),
    "minutes"
)

print(
    "Fitness:",
    round(average_fitness, 2),
    "minutes"
)

print(
    "Study:",
    round(average_study, 2),
    "minutes"
)

print(
    "Coding:",
    round(average_coding, 2),
    "minutes"
)

print(
    "Class:",
    round(average_class, 2),
    "minutes"
)

print(
    "Other Activities:",
    round(average_other, 2),
    "minutes"
)

print(
    "Free / Unaccounted:",
    round(average_free, 2),
    "minutes"
)


# ------------------------------------------------------------
# INDICES
# ------------------------------------------------------------

print()
print("INDEX RESULTS")
print("-" * 60)

print(
    "PAI:",
    round(PAI, 2)
)

print(
    "TPI:",
    round(TPI, 2)
)

print(
    "AAI:",
    round(AAI, 2)
)

print(
    "PhAI:",
    round(PhAI, 2)
)

print(
    "SRI:",
    round(SRI, 2)
)

print(
    "ABI:",
    round(ABI, 2)
)

print(
    "TUI:",
    round(TUI, 2)
)

print(
    "EI:",
    round(EI, 2)
)

print(
    "DCI:",
    round(DCI, 2),
    "%"
)


# ------------------------------------------------------------
# RELATIONSHIPS
# ------------------------------------------------------------

print()
print("RELATIONSHIP ANALYSIS")
print("-" * 60)

print()
print(
    "Sleep ↔ Energy:",
    round(sleep_energy, 3)
)

print(
    interpretation(sleep_energy)
)


print()
print(
    "Study ↔ Satisfaction:",
    round(study_satisfaction, 3)
)

print(
    interpretation(study_satisfaction)
)


print()
print(
    "Coding ↔ Energy:",
    round(coding_energy, 3)
)

print(
    interpretation(coding_energy)
)


# ------------------------------------------------------------
# FINISHED
# ------------------------------------------------------------

print()
print("=" * 60)

print("PROJECT COMPLETED SUCCESSFULLY")

print("=" * 60)
